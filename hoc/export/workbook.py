"""Excel view of hoc.db.

This is a generated view and is never read back: the database is the source of
truth (CLAUDE.md). Plain cells only, no merged ranges — the openpyxl merge and
row-insert pitfalls that bit the old workbook do not apply to a file that is
written fresh every time.
"""

from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.writer.excel import ExcelWriter

from hoc.db import HOUSE_BLOCK_FIELDS

DEFAULT_OUT_DIR = Path(__file__).resolve().parent.parent.parent / "outputs"
WORKBOOK_NAME = "Riding_Tracker.xlsx"

# openpyxl stamps a new Workbook's created/modified properties with the
# wall-clock now() by default, which would make two exports of the identical
# database differ by nothing but a timestamp — the one kind of churn this
# exporter otherwise avoids everywhere else (deterministic ordering, no
# generated-at line beyond the site's single turn/season marker). Pinned to a
# fixed sentinel so re-running the exporter against unchanged data produces a
# byte-identical file, which is what lets the engine workflow's "nothing
# changed" check (a plain git diff) actually detect nothing changing.
_SENTINEL_TIMESTAMP = datetime(1970, 1, 1)

# The zip container format's own date field, separately from the XML content
# above: DOS dates cannot represent anything before 1980, so this cannot reuse
# _SENTINEL_TIMESTAMP.
_SENTINEL_ZIP_DATE = (1980, 1, 1, 0, 0, 0)


class _DeterministicZipFile(ZipFile):
    """A ZipFile whose entries never carry a wall-clock timestamp.

    openpyxl writes the parts of an xlsx two different ways, and both default
    to the wall-clock time when not told otherwise:

    - Most parts (styles.xml, docProps/app.xml, and so on) go through
      `archive.writestr(name, data)` with a plain string name; zipfile builds
      a ZipInfo for that with `date_time=time.localtime(time.time())`.
    - Each worksheet is written to a temp file first, then added with
      `archive.write(path, arcname)`; zipfile's `write()` takes the timestamp
      from the temp file's own mtime, which is just as much "now" — and by
      the time `write()` returns, that timestamp is already serialised into
      the entry's local file header, so fixing it up on the ZipInfo object
      afterward would leave the local header and the central directory
      disagreeing.

    Overriding both to force the same fixed date is what makes the resulting
    file byte-identical across two exports of the same data — the property
    fix above is necessary but not sufficient on its own.
    """

    def writestr(self, zinfo_or_arcname, data, *args, **kwargs):
        if isinstance(zinfo_or_arcname, str):
            zinfo = ZipInfo(zinfo_or_arcname, date_time=_SENTINEL_ZIP_DATE)
            zinfo.compress_type = self.compression
            zinfo_or_arcname = zinfo
        return super().writestr(zinfo_or_arcname, data, *args, **kwargs)

    def write(self, filename, arcname=None, compress_type=None, compresslevel=None):
        zinfo = ZipInfo(arcname or str(filename), date_time=_SENTINEL_ZIP_DATE)
        zinfo.compress_type = self.compression if compress_type is None else compress_type
        with open(filename, "rb") as f:
            data = f.read()
        return self.writestr(zinfo, data, compresslevel=compresslevel)


def _save_deterministically(wb, path):
    """Save `wb` so that re-running the exporter against unchanged data
    produces a byte-identical file — no timestamp anywhere in the archive.

    Two separate pieces of openpyxl behaviour would otherwise defeat this:
    Workbook.save() calls openpyxl.writer.excel.save_workbook(), whose first
    line unconditionally re-stamps `properties.modified` to `datetime.now()`
    regardless of what it was set to beforehand, so this calls the same
    ExcelWriter that save_workbook() uses directly, skipping only that one
    line; and _DeterministicZipFile (above) covers every archive member's own
    timestamp, which save_workbook() never touches at all.
    """
    archive = _DeterministicZipFile(path, "w", ZIP_DEFLATED, allowZip64=True)
    ExcelWriter(wb, archive).save()

__all__ = ["write_workbook", "DEFAULT_OUT_DIR", "WORKBOOK_NAME"]

HEADER_FONT = Font(bold=True)


def _block_sort_key(field):
    """House-block headings in reading order; newly recovered ones after them."""
    if field in HOUSE_BLOCK_FIELDS:
        return (HOUSE_BLOCK_FIELDS.index(field), "")
    return (len(HOUSE_BLOCK_FIELDS), field)


def _sheet(wb, title, headers, rows, hex_columns=()):
    """Add a sheet with a bold, frozen header row.

    `hex_columns` are 1-based column indexes whose cells are filled with the
    colour they name, so the palette is visible in the workbook itself.
    """
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append(list(row))

    for column_index in hex_columns:
        for row_index in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_index, column=column_index)
            value = cell.value
            if isinstance(value, str) and value.startswith("#") and len(value) == 7:
                cell.fill = PatternFill("solid", fgColor=value[1:])

    for column_index, header in enumerate(headers, start=1):
        width = max(len(str(header)) + 2, 12)
        for row in rows:
            value = row[column_index - 1]
            if value is not None:
                width = max(width, min(len(str(value)) + 2, 60))
        ws.column_dimensions[ws.cell(row=1, column=column_index).column_letter].width = width
    return ws


def _riding_tracker_rows(conn):
    return conn.execute(
        "SELECT hs.house, hd.name AS holder, hd.generation, r.name_en AS riding, r.province,"
        "       hs.rank, hs.peerage, h.hex, h.seat_order"
        " FROM holdings h"
        " JOIN houses hs ON hs.house = h.house"
        " JOIN ridings r ON r.fed_id = h.fed_id"
        " LEFT JOIN holders hd ON hd.house = h.house AND hd.is_current = 1"
        " WHERE h.released_event_id IS NULL"
        " ORDER BY hs.house, h.seat_order"
    ).fetchall()


def _house_rows(conn):
    return conn.execute(
        "SELECT c.house, hs.status, hs.rank, hs.peerage, c.primary_hex, c.secondary_hex,"
        "       hd.name AS holder, hd.generation, hd.heir_apparent,"
        "       cl.personal_year, cl.basis,"
        "       (SELECT COUNT(*) FROM holdings h"
        "         WHERE h.house = c.house AND h.released_event_id IS NULL) AS ridings"
        " FROM v_house_colours c"
        " JOIN houses hs ON hs.house = c.house"
        " LEFT JOIN holders hd ON hd.house = c.house AND hd.is_current = 1"
        " LEFT JOIN clocks cl ON cl.house = c.house"
        " ORDER BY hs.status, c.house"
    ).fetchall()


def _matrix(conn):
    """Square house x house grid of relation markers, active houses only."""
    houses = [
        row["house"]
        for row in conn.execute("SELECT house FROM houses WHERE status = 'active' ORDER BY house")
    ]
    index = {house: i for i, house in enumerate(houses)}
    grid = [["" for _ in houses] for _ in houses]

    for row in conn.execute("SELECT house_a, house_b, marker FROM relations ORDER BY id"):
        a, b, marker = row["house_a"], row["house_b"], row["marker"]
        if a not in index or b not in index or not marker:
            continue  # removed houses appear in the long form, not the grid
        for i, j in ((index[a], index[b]), (index[b], index[a])):
            existing = grid[i][j]
            if not existing:
                grid[i][j] = marker
            elif marker not in existing.split("; "):
                grid[i][j] = f"{existing}; {marker}"

    headers = ["House"] + houses
    rows = [[house] + grid[i] for i, house in enumerate(houses)]
    return headers, rows


def write_workbook(conn, out_dir=DEFAULT_OUT_DIR):
    """Write outputs/Riding_Tracker.xlsx. Returns the path."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / WORKBOOK_NAME

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.created = _SENTINEL_TIMESTAMP
    wb.properties.modified = _SENTINEL_TIMESTAMP

    _sheet(
        wb,
        "Riding Tracker",
        ["House", "Holder", "Generation", "Riding", "Province", "Rank", "Peerage", "Hex", "Seat Order"],
        [tuple(row) for row in _riding_tracker_rows(conn)],
        hex_columns=(8,),
    )

    _sheet(
        wb,
        "Houses",
        ["House", "Status", "Rank", "Peerage", "Primary", "Secondary", "Holder", "Generation",
         "Heir apparent", "Personal year", "Clock basis", "Ridings"],
        [tuple(row) for row in _house_rows(conn)],
        hex_columns=(5, 6),
    )

    _sheet(
        wb,
        "House Blocks",
        ["House", "Field", "Text", "Source"],
        [
            (r["house"], r["field"], r["text"], r["source"])
            for r in sorted(
                conn.execute("SELECT * FROM house_blocks"),
                key=lambda r: (r["house"], _block_sort_key(r["field"])),
            )
        ],
    )

    _sheet(
        wb,
        "Successions",
        ["Seq", "House", "Predecessor", "Successor", "Transition", "Personal date", "Nature",
         "Batch", "Source"],
        [tuple(row) for row in conn.execute("SELECT * FROM successions ORDER BY seq")],
    )

    _sheet(
        wb,
        "Climate",
        ["Era cohort", "Seq", "Event", "Magnitude", "Tag", "Cumulative after", "Source"],
        [
            (r["era_cohort"], r["seq"], r["event"], r["magnitude"], r["tag"], r["cumulative_after"], r["source"])
            for r in conn.execute("SELECT * FROM climate ORDER BY era_cohort, seq")
        ],
    )

    _sheet(
        wb,
        "Relations",
        ["House A", "House B", "Marker", "Event", "Event id", "Source"],
        [
            (r["house_a"], r["house_b"], r["marker"], r["event_text"], r["event_id"], r["source"])
            for r in conn.execute("SELECT * FROM relations ORDER BY house_a, house_b, id")
        ],
    )

    matrix_headers, matrix_rows = _matrix(conn)
    _sheet(wb, "Matrix", matrix_headers, matrix_rows)

    _sheet(
        wb,
        "Events",
        ["Id", "Turn", "Kind", "Era cohort", "Title", "Houses", "Source", "Created", "Narrative"],
        [
            (
                r["id"],
                r["turn_id"],
                r["kind"],
                r["era_cohort"],
                r["title"],
                ", ".join(
                    h["house"]
                    for h in conn.execute(
                        "SELECT house FROM event_houses WHERE event_id = ? ORDER BY house", (r["id"],)
                    )
                ),
                r["source"],
                r["created_at"],
                r["narrative"],
            )
            for r in conn.execute("SELECT * FROM events ORDER BY id")
        ],
    )

    _sheet(
        wb,
        "Watch",
        ["Id", "Item", "Status"],
        [(r["id"], r["text"], r["status"]) for r in conn.execute("SELECT * FROM watch ORDER BY id")],
    )

    _sheet(
        wb,
        "Handoff",
        ["Key", "Text"],
        [(r["key"], r["text"]) for r in conn.execute("SELECT * FROM handoff ORDER BY key")],
    )

    _save_deterministically(wb, path)
    return path
