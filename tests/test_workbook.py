"""The Excel exporter must be byte-for-byte deterministic.

This is the one property the engine workflow's "nothing changed, don't commit"
check depends on: `rebuild` re-exports every output from an unchanged database,
and `git diff --cached --quiet` can only see "nothing changed" if a rebuild's
outputs/Riding_Tracker.xlsx really is nothing changed. openpyxl's own
Workbook.save() stamps two independent wall-clock timestamps into the file on
every save — the document properties, and every zip entry's own date field —
so two exports of the identical database used to differ every single time,
which would have made every `rebuild` run spuriously commit a no-op diff.
"""

import time
from pathlib import Path

from hoc import scenario
from hoc.export import workbook

ROOT = Path(__file__).resolve().parent.parent


def _load_seed_module():
    import importlib.util

    spec = importlib.util.spec_from_file_location("load_seed", ROOT / "scripts" / "load_seed.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_two_exports_of_the_same_database_are_byte_identical(tmp_path):
    conn = _load_seed_module().build(
        tmp_path / "hoc.db", seed=scenario.seed_dir("legacy")
    )

    first = workbook.write_workbook(conn, out_dir=tmp_path / "one")
    # A real gap, not just a different process tick: openpyxl's nondeterminism
    # here is measured in whole seconds (DOS zip dates have 2-second
    # resolution), so a sub-second gap could pass by accident.
    time.sleep(2)
    second = workbook.write_workbook(conn, out_dir=tmp_path / "two")
    conn.close()

    assert first.read_bytes() == second.read_bytes()


def test_the_workbook_still_opens_and_reads_correctly(tmp_path):
    """Determinism must not come at the cost of a workbook Excel or openpyxl
    itself can no longer read."""
    from openpyxl import load_workbook

    conn = _load_seed_module().build(
        tmp_path / "hoc.db", seed=scenario.seed_dir("legacy")
    )
    path = workbook.write_workbook(conn, out_dir=tmp_path)
    conn.close()

    wb = load_workbook(path)
    assert "Riding Tracker" in wb.sheetnames
    sheet = wb["Riding Tracker"]
    assert sheet.max_row > 1
    assert sheet.cell(row=1, column=1).value == "House"
