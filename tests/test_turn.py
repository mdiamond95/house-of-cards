"""Turn runner and exporter tests.

Every test rebuilds the database from the seed into a temporary path and writes
outputs into a temporary directory, so nothing here touches the committed
hoc.db or outputs/.
"""

import importlib.util
import json
from pathlib import Path

import pytest

from hoc.export import dump, map as map_export, workbook
from hoc.turn import TurnError, apply_turn

ROOT = Path(__file__).resolve().parent.parent


def _load_seed_module():
    spec = importlib.util.spec_from_file_location("load_seed", ROOT / "scripts" / "load_seed.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def conn(tmp_path):
    connection = _load_seed_module().build(tmp_path / "hoc.db")
    yield connection
    connection.close()


def narrative(words):
    """A narrative of a given length — the runner counts words, not prose."""
    return " ".join(["Nipigon"] * words)


def write_turn(tmp_path, name, data):
    turns_dir = tmp_path / "turns"
    turns_dir.mkdir(exist_ok=True)
    path = turns_dir / name
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return path


def unclaimed_land_neighbour(conn, house):
    """An unheld riding sharing a land border with a current holding of `house`."""
    row = conn.execute(
        "SELECT r.name_en FROM ridings r"
        " JOIN adjacency a ON (a.fed_id_a = r.fed_id OR a.fed_id_b = r.fed_id)"
        "   AND a.adjacency_type = 'land'"
        " JOIN holdings h ON h.released_event_id IS NULL AND h.house = ?"
        "   AND h.fed_id = CASE WHEN a.fed_id_a = r.fed_id THEN a.fed_id_b ELSE a.fed_id_a END"
        " WHERE r.fed_id NOT IN (SELECT fed_id FROM holdings WHERE released_event_id IS NULL)"
        " ORDER BY r.fed_id LIMIT 1",
        (house,),
    ).fetchone()
    assert row is not None, f"no unclaimed land neighbour for {house}"
    return row["name_en"]


def holdings_count(conn):
    return conn.execute("SELECT COUNT(*) FROM holdings WHERE released_event_id IS NULL").fetchone()[0]


# -------------------------------------------------------------- applying --


def test_valid_turn_applies_and_exports(conn, tmp_path):
    riding = unclaimed_land_neighbour(conn, "Whitcombe")
    path = write_turn(
        tmp_path,
        "0001_whitcombe-expansion.json",
        {
            "directive": "Whitcombe presses north-west along the timber line.",
            "era_cohort": "founding-era",
            "event": {
                "kind": "expansion",
                "title": "The Nipigon Concession",
                "houses": ["Whitcombe"],
                "narrative": narrative(120),
            },
            "operations": [{"op": "expand", "house": "Whitcombe", "riding": riding}],
            "watch": {"add": ["Does the concession bind the successor?"]},
            "threads": ["Whitcombe timber financing"],
        },
    )

    summary = apply_turn(conn, path)

    assert summary["turn_id"] == 1
    assert summary["ops_applied"] == 1
    assert summary["holdings"] == 144
    assert holdings_count(conn) == 144
    assert summary["warnings"] == []

    turn_row = conn.execute("SELECT * FROM turns WHERE turn_id = 1").fetchone()
    assert turn_row["directive"].startswith("Whitcombe presses")
    assert turn_row["status"] == "applied"

    event = conn.execute("SELECT * FROM events WHERE id = ?", (summary["event_id"],)).fetchone()
    assert event["kind"] == "expansion"
    assert event["turn_id"] == 1
    assert event["era_cohort"] == "founding-era"

    assert conn.execute("SELECT COUNT(*) FROM watch WHERE status = 'open'").fetchone()[0] == 1
    assert conn.execute("SELECT COUNT(*) FROM threads").fetchone()[0] == 1

    out_dir = tmp_path / "outputs"
    dump.write_dump(conn, out_dir=out_dir)
    state = json.loads((out_dir / "dump" / "state.json").read_text(encoding="utf-8"))
    assert state["counts"]["ridings_claimed"] == 144
    whitcombe = next(h for h in state["houses"] if h["house"] == "Whitcombe")
    assert whitcombe["riding_count"] == 11


def test_failing_operation_rolls_the_whole_turn_back(conn, tmp_path):
    riding = unclaimed_land_neighbour(conn, "Whitcombe")
    path = write_turn(
        tmp_path,
        "0002_two-expansions.json",
        {
            "directive": "Whitcombe overreaches.",
            "era_cohort": None,
            "event": {
                "kind": "expansion",
                "title": "The overreach",
                "houses": ["Whitcombe"],
                "narrative": narrative(80),
            },
            "operations": [
                {"op": "expand", "house": "Whitcombe", "riding": riding},
                # Sudbury is Whitcombe's own; an occupied riding cannot be taken.
                {"op": "expand", "house": "Whitcombe", "riding": "Sudbury"},
            ],
        },
    )

    with pytest.raises(TurnError) as excinfo:
        apply_turn(conn, path)

    assert excinfo.value.operation_index == 1
    assert "occupied by Whitcombe" in excinfo.value.reason

    assert holdings_count(conn) == 143
    assert conn.execute("SELECT COUNT(*) FROM turns").fetchone()[0] == 0
    assert conn.execute("SELECT COUNT(*) FROM events WHERE turn_id IS NOT NULL").fetchone()[0] == 0


def test_narrative_over_the_hard_limit_is_rejected(conn, tmp_path):
    path = write_turn(
        tmp_path,
        "0003_too-long.json",
        {
            "directive": "Too many words.",
            "era_cohort": None,
            "event": {"kind": "other", "title": "Long", "houses": [], "narrative": narrative(601)},
            "operations": [],
        },
    )

    with pytest.raises(TurnError) as excinfo:
        apply_turn(conn, path)
    assert "601 words" in excinfo.value.reason
    assert conn.execute("SELECT COUNT(*) FROM turns").fetchone()[0] == 0


# ------------------------------------------------------------- founding --


def founding_turn(conn, acknowledge):
    riding = conn.execute(
        "SELECT name_en FROM ridings"
        " WHERE fed_id NOT IN (SELECT fed_id FROM holdings WHERE released_event_id IS NULL)"
        " ORDER BY fed_id LIMIT 1"
    ).fetchone()["name_en"]
    operation = {
        "op": "found_house",
        "house": "Testerly",
        "peerage": "Baron Testerly of Testing",
        "rank": "Baron",
        "riding": riding,
        "primary_hex": "#123456",
        "secondary_hex": None,
        "holder_name": "Test Holder",
        "bio_age": "~50",
        "heir_apparent": None,
        "tag": "Progressive",
    }
    if acknowledge:
        operation["acknowledge_cohort_mismatch"] = True
    return {
        "directive": "Found a progressive house in a conservative season.",
        "era_cohort": "founding-era",
        "event": {
            "kind": "founding",
            "title": "A grant against the grain",
            "houses": [],
            "narrative": narrative(100),
        },
        "operations": [operation],
    }


def test_founding_against_the_climate_is_blocked(conn, tmp_path):
    # The founding-era ledger stands at -1, so a Progressive cohort mismatches.
    path = write_turn(tmp_path, "0004_founding.json", founding_turn(conn, acknowledge=False))

    with pytest.raises(TurnError) as excinfo:
        apply_turn(conn, path)

    assert excinfo.value.operation_index == 0
    assert "mismatch" in excinfo.value.reason
    assert conn.execute("SELECT COUNT(*) FROM houses WHERE house = 'Testerly'").fetchone()[0] == 0
    assert conn.execute("SELECT COUNT(*) FROM turns").fetchone()[0] == 0


def test_founding_proceeds_when_the_mismatch_is_acknowledged(conn, tmp_path):
    path = write_turn(tmp_path, "0005_founding-ack.json", founding_turn(conn, acknowledge=True))

    summary = apply_turn(conn, path)

    house = conn.execute("SELECT * FROM houses WHERE house = 'Testerly'").fetchone()
    assert house["status"] == "active"
    assert house["primary_hex"] == "#123456"

    holder = conn.execute("SELECT * FROM holders WHERE house = 'Testerly'").fetchone()
    assert holder["generation"] == "G1"
    assert holder["is_current"] == 1

    holding = conn.execute(
        "SELECT * FROM holdings WHERE house = 'Testerly' AND released_event_id IS NULL"
    ).fetchone()
    assert holding["seat_order"] == 1
    assert holding["hex"] == "#123456"

    clock = conn.execute("SELECT * FROM clocks WHERE house = 'Testerly'").fetchone()
    assert clock["personal_year"] == 1867
    assert clock["basis"] == "founding grant"

    # The house links itself to the turn event, since it did not exist when the
    # event was recorded.
    role = conn.execute(
        "SELECT role FROM event_houses WHERE event_id = ? AND house = 'Testerly'",
        (summary["event_id"],),
    ).fetchone()
    assert role["role"] == "founder"

    delta = json.loads(
        conn.execute("SELECT mechanical_delta FROM events WHERE id = ?", (summary["event_id"],))
        .fetchone()["mechanical_delta"]
    )
    assert delta["founded"][0]["house"] == "Testerly"
    assert delta["cohort_mismatch_acknowledged"][0] == {
        "house": "Testerly",
        "tag": "Progressive",
        "era_cohort": "founding-era",
        "climate": -1,
    }


# ------------------------------------------------------------ exporters --


def test_exporters_run_on_the_loaded_database(conn, tmp_path):
    out_dir = tmp_path / "outputs"

    written = dump.write_dump(conn, out_dir=out_dir)
    assert (out_dir / "dump" / "state.json") in written
    assert (out_dir / "dump" / "houses.csv").exists()
    assert (out_dir / "dump" / "adjacency.csv").exists()

    book = workbook.write_workbook(conn, out_dir=out_dir)
    assert book.exists() and book.stat().st_size > 0

    from openpyxl import load_workbook

    sheets = load_workbook(book).sheetnames
    assert sheets == [
        "Riding Tracker", "Houses", "House Blocks", "Successions", "Climate", "Relations",
        "Matrix", "Events", "Watch", "Handoff",
    ]

    maps = map_export.write_maps(conn, out_dir=out_dir)
    assert [p.name for p in maps] == ["map.svg", "map_secondary.svg"]
    for path in maps:
        assert path.stat().st_size < map_export.MAX_BYTES
        text = path.read_text(encoding="utf-8")
        assert text.startswith("<svg") and text.rstrip().endswith("</svg>")
        assert text.count("<path") == 343


def test_secondary_map_distinguishes_the_principal_seat(conn, tmp_path):
    fills_primary, _ = map_export._fills(conn, use_secondary=False)
    fills_secondary, _ = map_export._fills(conn, use_secondary=True)

    seat = conn.execute(
        "SELECT fed_id FROM holdings WHERE house = 'Whitcombe' AND seat_order = 1"
        " AND released_event_id IS NULL"
    ).fetchone()["fed_id"]
    other = conn.execute(
        "SELECT fed_id FROM holdings WHERE house = 'Whitcombe' AND seat_order = 2"
        " AND released_event_id IS NULL"
    ).fetchone()["fed_id"]

    assert fills_primary[seat] == fills_primary[other] == "#2E5E3A"
    assert fills_secondary[seat] == "#2E5E3A"
    assert fills_secondary[other] == "#4A7A55"
